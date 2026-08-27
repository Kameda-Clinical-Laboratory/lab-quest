"""Build Word files for the cover and the full wrap (back + spine + front)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parent


def _zero_margins(section, width_mm: float, height_mm: float) -> None:
    section.page_width = Mm(width_mm)
    section.page_height = Mm(height_mm)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)


def _full_page_image(doc: Document, png: Path, width_mm: float, height_mm: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    pPr = p._p.get_or_add_pPr()
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:after'), '0')
    spacing.set(qn('w:line'), '240')
    spacing.set(qn('w:lineRule'), 'auto')
    p.add_run().add_picture(str(png), width=Mm(width_mm), height=Mm(height_mm))


def build(png: Path, out: Path, width_mm: float, height_mm: float) -> None:
    doc = Document()
    _zero_margins(doc.sections[0], width_mm, height_mm)
    _full_page_image(doc, png, width_mm, height_mm)
    doc.save(out)
    print('wrote', out)


def main() -> None:
    build(ROOT / 'cover-a4.png', ROOT / '冒険の書-表紙.docx', 210, 297)
    wrap = ROOT / 'wrap-a4-12mm.png'
    if wrap.exists():
        build(wrap, ROOT / '冒険の書-カバー巻き.docx', 432, 297)


if __name__ == '__main__':
    main()
